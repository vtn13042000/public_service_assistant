"""scripts/export_slides.py: PowerPoint sections come from deck.json "sections" (start slide → next start)."""
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
pytest.importorskip("pptx")
import export_slides as X  # noqa: E402

SLIDE = ('<section id="{id}" style="background:#FBFBF8;color:#14213D;padding:128px;display:flex;flex-direction:column;gap:24px">'
         '<h2 style="font-size:72px">{id}</h2><p style="font-size:28px">body</p><aside>notes {id}</aside></section>')


def test_sections_follow_deck_json(tmp_path):
    order = ["cover", "a", "b", "c1", "c2"]
    (tmp_path / "slides").mkdir()
    for sid in order:
        (tmp_path / "slides" / f"{sid}.html").write_text(SLIDE.format(id=sid), encoding="utf-8")
    sections = {"s1": {"description": "Intro", "start": "cover"}, "s2": {"description": "Body", "start": "a"},
                "s3": {"description": "Appendix", "start": "c1"}}
    (tmp_path / "deck.json").write_text(json.dumps({"title": "T", "order": order, "sections": sections}), encoding="utf-8")
    out = tmp_path / "deck.pptx"
    X.export(tmp_path, out, fallback=False)

    from pptx import Presentation
    prs = Presentation(out)
    ids = [s.slide_id for s in prs.slides]
    secs = prs.part._element.findall(f".//{{{X.P14}}}section")
    got = [(s.get("name"), [int(i.get("id")) for i in s.iter(f"{{{X.P14}}}sldId")]) for s in secs]
    assert got == [("Intro", ids[0:1]), ("Body", ids[1:3]), ("Appendix", ids[3:5])]
    assert prs.slides[0].notes_slide.notes_text_frame.text == "notes cover"


def test_source_footer_dropped_unless_kept(tmp_path):
    slide = ('<section id="x" style="background:#FBFBF8;color:#14213D;padding:128px;display:flex;flex-direction:column;gap:24px">'
             '<h2 style="font-size:72px">x</h2><p style="font-size:28px">body</p>'
             '<p style="position:absolute;left:128px;bottom:64px;font-size:24px">Source: docs/a.md</p>'
             '<p style="position:absolute;right:128px;top:64px;font-size:24px">pinned label</p><aside>n</aside></section>')
    (tmp_path / "slides").mkdir()
    (tmp_path / "slides" / "x.html").write_text(slide, encoding="utf-8")
    (tmp_path / "deck.json").write_text(json.dumps({"title": "T", "order": ["x"]}), encoding="utf-8")

    from pptx import Presentation

    def texts(keep):
        out = tmp_path / f"deck-{keep}.pptx"
        X.export(tmp_path, out, fallback=False, keep_sources=keep)
        shapes = Presentation(out).slides[0].shapes
        return {sh.text_frame.text: (sh.left, sh.top) for sh in shapes if sh.has_text_frame and sh.text_frame.text}

    dropped, kept = texts(False), texts(True)
    assert not any(t.startswith("Source:") for t in dropped)
    assert "Source: docs/a.md" in kept and "pinned label" in dropped
    assert dropped == {t: pos for t, pos in kept.items() if not t.startswith("Source:")}


def test_blob_image_resolved_through_assets_json_and_contained(tmp_path):
    from PIL import Image

    Image.new("RGB", (400, 100), "red").save(tmp_path / "wide.png")  # 4:1 picture in a 400×400 box → 400×100, centred
    deck = tmp_path / "deck"
    (deck / "slides").mkdir(parents=True)
    (tmp_path / "assets.json").write_text(json.dumps({"/_blob/abc": "wide.png"}), encoding="utf-8")  # one level above the deck
    slide = ('<section id="x" style="background:#FBFBF8;padding:128px;display:flex;flex-direction:column;gap:20px">'
             '<h2 style="font-size:64px">x</h2><div style="display:flex;gap:24px;align-items:flex-start">'
             '<img src="/_blob/abc" style="width:400px;height:400px;object-fit:contain;border:1px solid #D9DEE8"></div></section>')
    (deck / "slides" / "x.html").write_text(slide, encoding="utf-8")
    (deck / "deck.json").write_text(json.dumps({"title": "T", "order": ["x"]}), encoding="utf-8")
    out = tmp_path / "deck.pptx"
    (sid, s, over), = X.export(deck, out, fallback=False)
    assert over == 0

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    pics = [sh for sh in Presentation(out).slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1
    p = pics[0]
    assert abs(p.width / p.height - 4.0) < 0.01
    assert abs(p.width - X.E(398)) <= X.E(1)  # 400 px box minus the 1 px border on each side

    (deck / "slides" / "x.html").write_text(slide.replace("/_blob/abc", "/_blob/missing"), encoding="utf-8")
    with pytest.raises(SystemExit):
        X.export(deck, tmp_path / "bad.pptx", fallback=False)
