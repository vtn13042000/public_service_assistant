"""Export the overview deck (docs/10_slides/project/) to a static PowerPoint and PDF.

    uv run python scripts/export_slides.py                    # → docs/10_slides/overview.pptx + overview.pdf
    uv run python scripts/export_slides.py --png /tmp/check   # also render every PDF page to PNG for a visual check
    uv run python scripts/export_slides.py --fallback-fonts --no-pdf --out docs/10_slides/overview-fallback-fonts.pptx
                                                              # Georgia / Arial instead of Source Serif 4 / IBM Plex Sans

Input: `deck.json` (slide order, title) and one HTML file per slide, written in a closed subset: a <section> canvas
(1920×1080 px), <div> boxes with inline flex styles (display:flex, flex-direction, gap, flex:N, width px/%, padding,
background, border, border-radius, justify-content, align-items), h1/h2/h3/p/ul/li text with <b>/<i>/<br>, <table>
with tr/th/td, <img> with a px width/height (src = an artifact asset url /_blob/<id>, mapped to a local file by
assets.json; fitted like object-fit: contain), an absolutely positioned footer <p>, and an <aside> with the speaker notes. A pinned footer that
starts with "Source:" is left out of the static files unless --keep-sources is given (the HTML deck keeps it).

How: a small flexbox-like layout pass measures every text block with the real font files (Pillow + fontconfig), so
line wrapping matches what LibreOffice/PowerPoint draw. A text-only box (a "card") becomes one shape with its fill,
border and rounded corners; a table becomes a pptx table; the first h1/h2 goes into the slide's title placeholder; the
<aside> becomes the speaker notes. If a slide's content is taller than the canvas, font sizes and gaps are scaled down
in small steps until it fits (the scale is printed per slide). Static only: no transitions or animations.
The PDF is made by `soffice --headless --convert-to pdf`. Pure Python apart from that call; re-runnable.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

import lxml.html
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
CANVAS_W, CANVAS_H = 1920, 1080
EMU_PER_PX = 6350                 # 1920 px = 13.333 in = 12 192 000 EMU (16:9)
PT_PER_PX = 0.5                   # 1 px on the 1920 canvas = 0.5 pt on the slide
NORMAL_LH = 1.3                   # CSS line-height "normal" (close to both fonts' ascent+descent)
CELL_PAD = (12, 16)               # table cell padding (vertical, horizontal) in px
WRAP_SAFETY = 0.97                # measure lines against 97 % of the box width: renderers wrap a little earlier
MIN_SCALE = 0.6
INHERITED = ("color", "font-family", "font-size", "font-weight", "font-style", "line-height", "text-align",
             "letter-spacing", "text-transform")
GENERIC = ("serif", "sans-serif", "monospace")
USE_FALLBACK = False              # --fallback-fonts: lay out and write with the second family of each font-family list
ASSETS: dict[str, Path] = {}       # artifact asset url (/_blob/<id>) → local file, from assets.json next to or above the deck
KEEP_SOURCES = False              # --keep-sources: keep the pinned "Source: …" footers (the HTML deck always has them)
DEFAULT_SIZES = {"h1": 96, "h2": 72, "h3": 34, "p": 28, "li": 28, "td": 28, "th": 28}


def px(v: str | None, ref: float = 0.0) -> float:
    """'24px' → 24, '40%' → 0.4·ref, '0' → 0, None → 0."""
    if not v:
        return 0.0
    v = v.strip()
    if v.endswith("%"):
        return float(v[:-1]) / 100 * ref
    m = re.match(r"(-?[\d.]+)", v)
    return float(m.group(1)) if m else 0.0


def parse_style(s: str | None) -> dict:
    out = {}
    for part in (s or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = v.strip()
    return out


def box4(v: str | None) -> tuple[float, float, float, float]:
    """CSS padding shorthand → (top, right, bottom, left)."""
    p = [px(x) for x in (v or "0").split()]
    if len(p) == 1:
        return p[0], p[0], p[0], p[0]
    if len(p) == 2:
        return p[0], p[1], p[0], p[1]
    if len(p) == 3:
        return p[0], p[1], p[2], p[1]
    return p[0], p[1], p[2], p[3]


def rgb(v: str | None) -> RGBColor | None:
    if not v:
        return None
    m = re.search(r"#([0-9a-fA-F]{6}|[0-9a-fA-F]{3})\b", v)
    if not m:
        return None
    h = m.group(1)
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    return RGBColor.from_string(h.upper())


def families(v: str | None) -> list[str]:
    return [f.strip().strip("'\"") for f in (v or "").split(",") if f.strip()]


# ---------------------------------------------------------------- fonts and measuring

@lru_cache(maxsize=None)
def font_file(family: str, bold: bool, italic: bool) -> str:
    """The font file fontconfig would draw `family` with (the same substitution LibreOffice gets)."""
    pattern = family + (":bold" if bold else "") + (":italic" if italic else "")
    try:
        return subprocess.run(["fc-match", "-f", "%{file}", pattern], capture_output=True, text=True,
                              check=True).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return ""


@lru_cache(maxsize=None)
def pil_font(family: str, bold: bool, italic: bool, size10: int):
    path = font_file(family, bold, italic)
    try:
        return ImageFont.truetype(path, max(1, size10) / 10)
    except OSError:
        return ImageFont.load_default()


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    color: RGBColor | None = None
    size: float = 28.0            # px, before scaling
    family: str = "IBM Plex Sans"
    spacing: float = 0.0          # letter-spacing px
    br: bool = False              # a line break, no text

    def width(self, text: str, s: float) -> float:
        f = pil_font(self.family, self.bold, self.italic, round(self.size * s * 10))
        return f.getlength(text) + self.spacing * s * len(text)


def tokens(runs: list[Run]):
    """(run, word-with-trailing-space) pairs; a break run yields (run, None)."""
    for r in runs:
        if r.br:
            yield r, None
            continue
        for w in re.findall(r"\S+\s*|\s+", r.text):
            yield r, w


def wrap_lines(runs: list[Run], width: float, s: float) -> tuple[int, float]:
    """Greedy word wrap → (number of lines, widest line in px)."""
    lines, cur, widest = 1, 0.0, 0.0
    limit = width * WRAP_SAFETY
    for r, w in tokens(runs):
        if w is None:
            widest, lines, cur = max(widest, cur), lines + 1, 0.0
            continue
        ww = r.width(w, s)
        if cur > 0 and cur + r.width(w.rstrip(), s) > limit:
            widest, lines, cur = max(widest, cur), lines + 1, 0.0
            if not w.strip():
                continue
        cur += ww
    return lines, max(widest, cur)


# ---------------------------------------------------------------- the layout tree

@dataclass
class Text:
    el: object
    cs: dict
    runs: list[Run]
    bullet: str = ""
    absolute: dict | None = None
    flex: float = 0.0
    width_css: str | None = None
    is_title: bool = False
    st: dict = field(default_factory=dict)   # own box styles: a text element may carry background/padding (a "pill")

    @property
    def pad(self):
        return box4(self.st.get("padding"))

    @property
    def bg(self):
        return rgb(self.st.get("background") or self.st.get("background-color"))

    @property
    def size(self) -> float:
        return px(self.cs.get("font-size")) or 28

    def lh(self, s: float) -> float:
        v = self.cs.get("line-height", "")
        mult = float(v) if re.fullmatch(r"[\d.]+", v or "") else (px(v) / self.size if v.endswith("px") else NORMAL_LH)
        return self.size * s * mult

    def runs_s(self) -> list[Run]:
        return ([Run(self.bullet, family=self.runs[0].family, size=self.size)] if self.bullet and self.runs else []) + self.runs

    def height(self, w: float, s: float) -> float:
        t, r, b, l = self.pad
        return wrap_lines(self.runs_s(), w - l - r, s)[0] * self.lh(s) + t + b

    def intrinsic(self, s: float) -> float:
        t, r, b, l = self.pad
        return wrap_lines(self.runs_s(), 1e9, s)[1] / WRAP_SAFETY + 2 + l + r

    def min_content(self, s: float) -> float:
        """Width of the longest unbreakable word (CSS min-content): a flex item never shrinks below it."""
        longest = max((r.width(w.strip(), s) for r, w in tokens(self.runs_s()) if w and w.strip()), default=0)
        t, r, b, l = self.pad
        return longest * 1.06 + 6 + l + r  # slack: renderers must never break inside a word


@dataclass
class Table:
    el: object
    cs: dict
    rows: list[dict]              # {"bg": RGBColor|None, "cells": [Text], "header": bool}
    widths: list[float]           # fractions
    border: tuple[float, RGBColor | None] = (0, None)
    absolute: dict | None = None
    flex: float = 0.0
    width_css: str | None = None

    def row_heights(self, w: float, s: float) -> list[float]:
        out = []
        for row in self.rows:
            h = 0.0
            for c, f in zip(row["cells"], self.widths):
                h = max(h, c.height(w * f - 2 * CELL_PAD[1], s))
            out.append(h + 2 * CELL_PAD[0])
        return out

    def height(self, w: float, s: float) -> float:
        return sum(self.row_heights(w, s))

    def intrinsic(self, s: float) -> float:
        return CANVAS_W

    def min_content(self, s: float) -> float:
        return 0.0


@dataclass
class Image:
    """An <img> with a px width/height box; the picture is fitted inside it like object-fit: contain."""
    el: object
    path: Path
    st: dict
    absolute: dict | None = None
    flex: float = 0.0

    @property
    def width_css(self):
        return self.st.get("width")

    @property
    def border(self):
        b = self.st.get("border")
        return (px(b.split()[0]) if b else 0.0, rgb(b))

    def height(self, w: float, s: float) -> float:
        return px(self.st.get("height")) or w * self.ratio()

    def ratio(self) -> float:
        from PIL import Image as PILImage
        with PILImage.open(self.path) as im:
            return im.height / im.width

    def intrinsic(self, s: float) -> float:
        return px(self.st.get("width")) or CANVAS_W

    def min_content(self, s: float) -> float:
        return px(self.st.get("width"))


@dataclass
class Box:
    el: object
    cs: dict
    st: dict
    children: list = field(default_factory=list)

    @property
    def row(self) -> bool:
        return "flex" in self.st.get("display", "") and not self.st.get("flex-direction", "row").startswith("column")

    @property
    def pad(self):
        return box4(self.st.get("padding"))

    @property
    def flex(self) -> float:
        return px(self.st.get("flex", "").split()[0]) if self.st.get("flex") else 0.0

    @property
    def width_css(self):
        return self.st.get("width")

    @property
    def absolute(self):
        return self.st if self.st.get("position") == "absolute" else None

    @property
    def bg(self):
        return rgb(self.st.get("background") or self.st.get("background-color"))

    @property
    def border(self):
        b = self.st.get("border")
        return (px(b.split()[0]) if b else 0.0, rgb(b))

    def gap(self, s: float) -> float:
        return px(self.st.get("gap")) * s

    def flow(self):
        return [c for c in self.children if not c.absolute]

    def text_only(self) -> bool:
        kids = self.flow()
        return bool(kids) and all(isinstance(c, Text) for c in kids) and (not self.row or len(kids) == 1)

    def min_content(self, s: float) -> float:
        t, r, b, l = self.pad
        kids = self.flow()
        if not kids:
            return l + r
        ws = [px(c.width_css) if c.width_css and not c.width_css.endswith("%") else c.min_content(s) for c in kids]
        return l + r + (sum(ws) + self.gap(s) * (len(ws) - 1) if self.row else max(ws))

    def widths(self, inner_w: float, s: float) -> list[float]:
        kids, g = self.flow(), self.gap(s)
        free = inner_w - g * max(0, len(kids) - 1)
        base, grow = [], []
        for c in kids:
            if c.width_css:
                base.append(px(c.width_css, inner_w))
                grow.append(0.0)
            elif c.flex:
                base.append(0.0)
                grow.append(c.flex)
            else:
                base.append(c.intrinsic(s))
                grow.append(0.0)
        rest = free - sum(base)
        tot = sum(grow)
        if rest >= 0:
            return [b + (rest * g_ / tot if tot else 0) for b, g_ in zip(base, grow)]
        # CSS flex-shrink: 1 — content-sized items give up space in proportion to their size, never below min-content
        mins = [c.min_content(s) if not c.width_css else b for c, b in zip(kids, base)]
        out = list(base)
        for _ in range(len(kids)):
            over = sum(out) - free
            loose = [i for i in range(len(kids)) if out[i] > mins[i] + 0.5]
            if over <= 0.5 or not loose:
                break
            tot_b = sum(out[i] for i in loose)
            for i in loose:
                out[i] = max(mins[i], out[i] - over * out[i] / tot_b)
        return out

    def overflow_x(self, inner_w: float, s: float) -> bool:
        return self.row and sum(self.widths(inner_w, s)) + self.gap(s) * max(0, len(self.flow()) - 1) > inner_w + 1

    def height(self, w: float, s: float) -> float:
        t, r, b, l = self.pad
        inner = w - l - r
        kids = self.flow()
        if not kids:
            return t + b
        if self.row:
            return t + b + max(c.height(cw, s) for c, cw in zip(kids, self.widths(inner, s)))
        return t + b + sum(c.height(inner, s) for c in kids) + self.gap(s) * (len(kids) - 1)

    def intrinsic(self, s: float) -> float:
        t, r, b, l = self.pad
        kids = self.flow()
        if not kids:
            return l + r
        ws = [c.intrinsic(s) for c in kids]
        return l + r + (sum(ws) + self.gap(s) * (len(ws) - 1) if self.row else max(ws))


def computed(el, parent: dict) -> tuple[dict, dict]:
    st = parse_style(el.get("style"))
    cs = {k: parent[k] for k in INHERITED if k in parent}
    cs.update({k: v for k, v in st.items() if k in INHERITED})
    if el.tag in DEFAULT_SIZES and "font-size" not in st:
        cs["font-size"] = f"{DEFAULT_SIZES[el.tag]}px" if el.tag not in ("td", "th", "li") or "font-size" not in cs \
            else cs["font-size"]
    if el.tag in ("h1", "h2", "h3", "th", "b", "strong") and "font-weight" not in st:
        cs["font-weight"] = "700"
    return st, cs


def is_bold(cs: dict) -> bool:
    w = cs.get("font-weight", "400")
    return w == "bold" or (w.isdigit() and int(w) >= 600)


def make_run(text: str, cs: dict) -> Run:
    fam = families(cs.get("font-family")) or ["IBM Plex Sans"]
    if USE_FALLBACK and len(fam) > 1 and fam[1] not in GENERIC:
        fam = fam[1:]
    if cs.get("text-transform") == "uppercase":
        text = text.upper()
    return Run(text, bold=is_bold(cs), italic=cs.get("font-style") == "italic", color=rgb(cs.get("color")),
               size=px(cs.get("font-size")) or 28, family=fam[0], spacing=px(cs.get("letter-spacing")))


INLINE_TAGS = ("b", "strong", "i", "em", "span", "a", "br", "code")


def collect_runs(el, cs: dict) -> list[Run]:
    out = []
    if el.text:
        out.append(make_run(el.text, cs))
    for ch in el:
        if not isinstance(ch.tag, str):
            pass
        elif ch.tag == "br":
            out.append(Run("", br=True))
        else:
            _, ccs = computed(ch, cs)
            if ch.tag in ("i", "em"):
                ccs["font-style"] = "italic"
            out.extend(collect_runs(ch, ccs))
        if ch.tail:
            out.append(make_run(ch.tail, cs))
    # collapse whitespace like HTML does
    for r in out:
        r.text = re.sub(r"\s+", " ", r.text)
    if out and not out[0].br:
        out[0].text = out[0].text.lstrip()
    if out and not out[-1].br:
        out[-1].text = out[-1].text.rstrip()
    return [r for r in out if r.br or r.text]


def build(el, parent_cs: dict):
    st, cs = computed(el, parent_cs)
    tag = el.tag
    absolute = st if st.get("position") == "absolute" else None
    flex = px(st.get("flex", "").split()[0]) if st.get("flex") else 0.0
    if tag in ("h1", "h2", "h3", "h4", "p", "li"):
        if absolute and not KEEP_SOURCES and "".join(el.itertext()).strip().startswith("Source:"):
            return None  # pinned source footer: out of the flow, so dropping it moves nothing else
        return Text(el, cs, collect_runs(el, cs), absolute=absolute, flex=flex, width_css=st.get("width"), st=st)
    if tag in ("ul", "ol"):
        box = Box(el, cs, {**st, "display": "flex", "flex-direction": "column", "gap": st.get("gap", "8px")})
        for i, li in enumerate(x for x in el if x.tag == "li"):
            t = build(li, cs)
            t.bullet = "• " if tag == "ul" else f"{i + 1}. "
            box.children.append(t)
        return box
    if tag == "table":
        return build_table(el, st, cs, absolute, flex)
    if tag == "img":
        return Image(el, resolve_asset(el.get("src", "")), st, absolute, flex)
    if tag in ("div", "section"):
        box = Box(el, cs, st)
        # A div holding only inline text (e.g. <div><b>…</b></div>) is one paragraph, as in the browser.
        kids = [ch for ch in el if isinstance(ch.tag, str) and ch.tag != "aside"]
        if kids and all(ch.tag in INLINE_TAGS for ch in kids) and "".join(el.itertext()).strip():
            box.children.append(Text(el, cs, collect_runs(el, cs)))
            return box
        for ch in el:
            if isinstance(ch.tag, str) and ch.tag != "aside":
                node = build(ch, cs)
                if node is not None:
                    box.children.append(node)
        return box
    return None


def build_table(el, st, cs, absolute, flex) -> Table:
    rows, widths = [], None
    for tr in el.iter("tr"):
        tst, tcs = computed(tr, cs)
        cells, ws = [], []
        for c in tr:
            if c.tag in ("td", "th"):
                cst, ccs = computed(c, tcs)
                cells.append(Text(c, ccs, collect_runs(c, ccs)))
                ws.append(px(cst.get("width"), 1.0) if cst.get("width", "").endswith("%") else 0)
        rows.append({"bg": rgb(tst.get("background") or tst.get("background-color")), "cells": cells,
                     "header": any(c.tag == "th" for c in tr)})
        if widths is None and all(ws):
            widths = ws
    n = max(len(r["cells"]) for r in rows)
    widths = widths or [1 / n] * n
    tot = sum(widths)
    b = st.get("border")
    return Table(el, cs, rows, [w / tot for w in widths], (px(b.split()[0]) if b else 0, rgb(b)), absolute, flex,
                 st.get("width") if st.get("width") != "100%" else None)


# ---------------------------------------------------------------- writing pptx shapes

def E(v: float) -> Emu:
    return Emu(int(round(v * EMU_PER_PX)))


class Writer:
    def __init__(self, slide, s: float):
        self.slide, self.s = slide, s

    def fill_paragraph(self, p, t: Text, space_before: float):
        s = self.s
        align = t.cs.get("text-align")
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
        p.line_spacing = Pt(t.lh(s) * PT_PER_PX)
        p.space_before = Pt(space_before * PT_PER_PX)
        p.space_after = Pt(0)
        for r in t.runs_s():
            if r.br:
                p.add_line_break()
                continue
            run = p.add_run()
            run.text = r.text
            f = run.font
            f.size = Pt(r.size * s * PT_PER_PX)
            f.bold, f.italic = r.bold, r.italic
            f.name = r.family
            if r.color is not None:
                f.color.rgb = r.color
            if r.spacing:
                run._r.get_or_add_rPr().set("spc", str(int(r.spacing * s * PT_PER_PX * 100)))

    def text_frame(self, shape, texts: list[Text], gap: float, pad=(0, 0, 0, 0), anchor=MSO_ANCHOR.TOP):
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        t, r, b, l = pad
        tf.margin_top, tf.margin_right, tf.margin_bottom, tf.margin_left = E(t), E(r), E(b), E(l)
        for i, tx in enumerate(texts):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            self.fill_paragraph(p, tx, gap if i else 0)
        return tf

    def textbox(self, texts, x, y, w, h, gap=0.0):
        sh = self.slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
        self.text_frame(sh, texts, gap)
        return sh

    def rect(self, x, y, w, h, bg, border, radius, texts=None, pad=(0, 0, 0, 0), gap=0.0):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        sh = self.slide.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
        if radius:
            sh.adjustments[0] = min(0.5, radius / max(1.0, min(w, h)))
            # the roundRect preset's text area is already inset by 0.293 × the corner radius on every side
            geo = 0.29289 * min(radius, min(w, h) / 2)
            pad = tuple(max(0.0, v - geo) for v in pad)
        if bg is not None:
            sh.fill.solid()
            sh.fill.fore_color.rgb = bg
        else:
            sh.fill.background()
        bw, bc = border
        if bw and bc is not None:
            sh.line.color.rgb = bc
            sh.line.width = E(bw)
        else:
            sh.line.fill.background()
        style = sh._element.find(qn("p:style"))  # theme style: drop shadow + accent colours; everything is set explicitly
        if style is not None:
            sh._element.remove(style)
        self.text_frame(sh, texts or [], gap, pad)
        return sh

    def image(self, im: Image, x, y, w, h):
        bw = im.border[0]
        bg = rgb(im.st.get("background") or im.st.get("background-color"))
        if bg is not None or bw:
            self.rect(x, y, w, h, bg, im.border, px(im.st.get("border-radius")))
        iw, ih = w - 2 * bw, h - 2 * bw  # object-fit: contain inside the border, centred
        r = im.ratio()
        pw, ph = (iw, iw * r) if iw * r <= ih else (ih / r, ih)
        return self.slide.shapes.add_picture(str(im.path), E(x + bw + (iw - pw) / 2), E(y + bw + (ih - ph) / 2), E(pw), E(ph))

    def title(self, t: Text, x, y, w, h):
        sh = self.slide.shapes.title
        if sh is None:
            return self.textbox([t], x, y, w, h)
        sh.left, sh.top, sh.width, sh.height = E(x), E(y), E(w), E(h)
        sh.text_frame.clear()
        self.text_frame(sh, [t], 0)
        return sh

    def table(self, tb: Table, x, y, w, h):
        s = self.s
        rh = tb.row_heights(w, s)
        ncol = len(tb.widths)
        gs = self.slide.shapes.add_table(len(tb.rows), ncol, E(x), E(y), E(w), E(sum(rh)))
        tbl = gs.table
        # plain table: no banding / header style from the theme
        tblPr = tbl._tbl.tblPr
        for a in ("firstRow", "bandRow"):
            tblPr.set(a, "0")
        style = tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
        for j, f in enumerate(tb.widths):
            tbl.columns[j].width = E(w * f)
        for i, row in enumerate(tb.rows):
            tbl.rows[i].height = E(rh[i])
            for j in range(ncol):
                cell = tbl.cell(i, j)
                if row["bg"] is not None:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = row["bg"]
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.margin_top = cell.margin_bottom = E(CELL_PAD[0])
                cell.margin_left = cell.margin_right = E(CELL_PAD[1])
                cell.vertical_anchor = MSO_ANCHOR.TOP
                tf = cell.text_frame
                tf.word_wrap = True
                if j < len(row["cells"]):
                    self.fill_paragraph(tf.paragraphs[0], row["cells"][j], 0)
                self._cell_borders(cell, tb.border)
        return gs

    @staticmethod
    def _cell_borders(cell, border):
        bw, bc = border
        if not (bw and bc is not None):
            return
        tcPr = cell._tc.get_or_add_tcPr()
        for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = tcPr.find(qn(tag))
            if ln is None:
                ln = tcPr.makeelement(qn(tag), {})
                tcPr.insert(0, ln)  # line properties precede the cell fill in CT_TableCellProperties
            ln.set("w", str(int(bw * EMU_PER_PX)))
            for ch in list(ln):
                ln.remove(ch)
            sf = ln.makeelement(qn("a:solidFill"), {})
            clr = sf.makeelement(qn("a:srgbClr"), {"val": str(bc)})
            sf.append(clr)
            ln.append(sf)


# ---------------------------------------------------------------- placing

def place(node, x, y, w, h, wr: Writer, section: Box):
    s = wr.s
    if isinstance(node, Text):
        if node.bg is not None and not node.is_title:  # a pill: filled (rounded) shape around its own text
            return wr.rect(x, y, w, h, node.bg, (0, None), px(node.st.get("border-radius")), [node], node.pad)
        return wr.title(node, x, y, w, h) if node.is_title else wr.textbox([node], x, y, w, h)
    if isinstance(node, Table):
        return wr.table(node, x, y, w, h)
    if isinstance(node, Image):
        return wr.image(node, x, y, w, h)
    t, r, b, l = node.pad
    radius = px(node.st.get("border-radius"))
    has_shape = node.bg is not None or node.border[0]
    if node is not section and node.text_only() and not any(c.is_title for c in node.flow()):
        if has_shape:
            return wr.rect(x, y, w, h, node.bg, node.border, radius, node.flow(), node.pad, node.gap(s))
        return wr.textbox(node.flow(), x + l, y + t, w - l - r, h - t - b, node.gap(s))
    if has_shape and node is not section:
        wr.rect(x, y, w, h, node.bg, node.border, radius)
    ix, iy, iw, ih = x + l, y + t, w - l - r, h - t - b
    kids = node.flow()
    g = node.gap(s)
    if node.row:
        ws = node.widths(iw, s)
        align = node.st.get("align-items", "stretch")
        cx = ix
        for c, cw in zip(kids, ws):
            ch = ih if align == "stretch" else c.height(cw, s)
            cy = iy + ((ih - ch) / 2 if align == "center" else (ih - ch) if align == "flex-end" else 0)
            place(c, cx, cy, cw, ch, wr, section)
            cx += cw + g
    else:
        hs = [c.height(iw, s) for c in kids]
        extra = ih - sum(hs) - g * max(0, len(kids) - 1)
        grow = [getattr(c, "flex", 0) for c in kids]
        if extra > 0 and sum(grow):
            hs = [hh + extra * gr / sum(grow) for hh, gr in zip(hs, grow)]
            extra = 0
        just = node.st.get("justify-content", "flex-start")
        cy, step = iy, g
        if extra > 0 and just == "space-between" and len(kids) > 1:
            step = g + extra / (len(kids) - 1)
        elif extra > 0 and just == "center":
            cy += extra / 2
        elif extra > 0 and just == "flex-end":
            cy += extra
        for c, hh in zip(kids, hs):
            place(c, ix, cy, iw, hh, wr, section)
            cy += hh + step
    for c in node.children:
        if c.absolute:
            a = c.absolute
            left = px(a.get("left"))
            cw = px(a.get("width"), CANVAS_W) or (CANVAS_W - left - (px(a.get("right")) if a.get("right") else left))
            chh = c.height(cw, s)
            top = px(a.get("top")) if a.get("top") else CANVAS_H - px(a.get("bottom")) - chh
            place(c, left, top, cw, chh, wr, section)


def fits(sec: Box, s: float) -> tuple[bool, float]:
    t, r, b, l = sec.pad
    iw, ih = CANVAS_W - l - r, CANVAS_H - t - b
    need = sec.height(CANVAS_W, s) - t - b
    ok = need <= ih + 0.5 and not any_overflow_x(sec, CANVAS_W, s)
    return ok, need - ih


def any_overflow_x(node, w, s) -> bool:
    if not isinstance(node, Box):
        return False
    t, r, b, l = node.pad
    iw = w - l - r
    if node.overflow_x(iw, s):
        return True
    kids = node.flow()
    ws = node.widths(iw, s) if node.row else [iw] * len(kids)
    return any(any_overflow_x(c, cw, s) for c, cw in zip(kids, ws))


# ---------------------------------------------------------------- deck

def load_section(path: Path):
    doc = lxml.html.fragment_fromstring(path.read_text(encoding="utf-8"), create_parent="div")
    sec = doc.find(".//section")
    aside = sec.find(".//aside")
    notes = re.sub(r"\s+", " ", aside.text_content()).strip() if aside is not None else ""
    root = build(sec, {})
    st = root.st
    if "display" not in st:
        st["display"] = "flex"
        st["flex-direction"] = "column"
    for tag in ("h1", "h2"):
        for t in walk(root):
            if isinstance(t, Text) and t.el.tag == tag:
                t.is_title = True
                return root, notes
    return root, notes


def load_assets(deck_dir: Path) -> None:
    """assets.json ({"/_blob/<id>": "relative/path.png"}) in the deck folder or one of its parents (one map for several decks);
    paths are relative to the map's own folder."""
    ASSETS.clear()
    for d in [deck_dir, *deck_dir.resolve().parents][:4]:
        f = d / "assets.json"
        if f.exists():
            ASSETS.update({k: (f.parent / v).resolve() for k, v in json.loads(f.read_text(encoding="utf-8")).items()})
            return


def resolve_asset(src: str) -> Path:
    key = src[src.find("/_blob/"):] if "/_blob/" in src else src  # an absolute artifact url maps like its path
    path = ASSETS.get(key)
    if path is None or not path.exists():
        raise SystemExit(f"image {src!r}: not in assets.json or the file is missing ({path})")
    return path


def walk(node):
    yield node
    for c in getattr(node, "children", []):
        yield from walk(c)


P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTION_EXT_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"  # the PowerPoint 2010 extension that carries slide sections


def add_sections(prs, order: list[str], sections: dict) -> list[tuple[str, str, int]]:
    """PowerPoint sections from deck.json "sections" ({key: {description, start}}): each runs from its start slide to the next
    section's start. python-pptx has no API for them, so the p14:sectionLst extension is written into presentation.xml.
    PowerPoint shows them in the slide sorter; LibreOffice and the PDF ignore them. Returns [(name, start id, n slides)]."""
    import uuid
    starts = sorted((order.index(v["start"]), v.get("description") or k) for k, v in sections.items() if v.get("start") in order)
    if not starts:
        return []
    if starts[0][0] != 0:  # slides before the first named section still need a section
        starts.insert(0, (0, "Untitled"))
    slide_ids = [sld.get("id") for sld in prs.slides._sldIdLst]
    pres = prs.part._element
    ext_lst = pres.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = pres.makeelement(qn("p:extLst"), {})
        pres.append(ext_lst)  # extLst is the last child of p:presentation
    ext = ext_lst.makeelement(qn("p:ext"), {"uri": SECTION_EXT_URI})
    sec_lst = ext.makeelement(f"{{{P14}}}sectionLst", nsmap={"p14": P14})
    report = []
    for i, (first, name) in enumerate(starts):
        last = starts[i + 1][0] if i + 1 < len(starts) else len(order)
        sec = sec_lst.makeelement(f"{{{P14}}}section", {"name": name, "id": "{" + str(uuid.uuid4()).upper() + "}"})
        ids = sec.makeelement(f"{{{P14}}}sldIdLst", {})
        for sid in slide_ids[first:last]:
            ids.append(ids.makeelement(f"{{{P14}}}sldId", {"id": sid}))
        sec.append(ids)
        sec_lst.append(sec)
        report.append((name, order[first], last - first))
    ext.append(sec_lst)
    ext_lst.append(ext)
    return report


def export(deck_dir: Path, out: Path, fallback: bool, keep_sources: bool = False) -> list[tuple[str, float, float]]:
    global USE_FALLBACK, KEEP_SOURCES
    KEEP_SOURCES = keep_sources
    load_assets(deck_dir)
    USE_FALLBACK = fallback  # measured with the fallback fonts too (fontconfig's substitutes when they are not installed)
    deck = json.loads((deck_dir / "deck.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(CANVAS_W), E(CANVAS_H)
    prs.core_properties.title = deck.get("title", "")
    layout = prs.slide_layouts[5]  # "Title Only": keeps a real title per slide (outline, accessibility)
    report = []
    for sid in deck["order"]:
        sec, notes = load_section(deck_dir / "slides" / f"{sid}.html")
        s = 1.0
        ok, over = fits(sec, s)
        while not ok and s > MIN_SCALE:
            s = round(s - 0.02, 2)
            ok, over = fits(sec, s)
        slide = prs.slides.add_slide(layout)
        bg = sec.bg
        if bg is not None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = bg
        place(sec, 0, 0, CANVAS_W, CANVAS_H, Writer(slide, s), sec)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        report.append((sid, s, over if not ok else 0.0))
    for name, start, n in add_sections(prs, deck["order"], deck.get("sections") or {}):
        print(f"section      {start:12s} +{n:<3d} {name}")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return report


def to_pdf(pptx: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("soffice not found; the .pptx was written, the PDF was not")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
                   check=True, capture_output=True)
    return pptx.with_suffix(".pdf")


def to_png(pdf: Path, out_dir: Path, dpi: int = 60) -> list[Path]:
    if not shutil.which("pdftoppm"):
        print("pdftoppm not found; skipping PNG check")
        return []
    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), str(pdf), str(out_dir / "page")], check=True)
    return sorted(out_dir.glob("page*.png"))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--deck", type=Path, default=ROOT / "docs/10_slides/project")
    ap.add_argument("--out", type=Path, default=ROOT / "docs/10_slides/overview.pptx")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--png", type=Path, help="render every PDF page to PNG in this directory")
    ap.add_argument("--fallback-fonts", action="store_true", help="Georgia/Arial (the HTML fallbacks) in the .pptx")
    ap.add_argument("--keep-sources", action="store_true",
                    help='keep the pinned "Source: …" footers (omitted from the static files by default)')
    a = ap.parse_args(argv)
    report = export(a.deck, a.out, a.fallback_fonts, a.keep_sources)
    bad = False
    for sid, s, over in report:
        flag = f"  OVERFLOW by {over:.0f}px at min scale" if over > 0 else ""
        bad |= over > 0
        print(f"{sid:12s} scale {s:.2f}{flag}")
    print(f"wrote {a.out}")
    if not a.no_pdf:
        pdf = to_pdf(a.out)
        print(f"wrote {pdf}")
        if a.png:
            pages = to_png(pdf, a.png)
            print(f"rendered {len(pages)} pages to {a.png}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
